import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1]

# (key, screen, origin, id, en, note)
# origin: DESIGN = verbatim from prototype, ANDROID = adapted from an iOS-only string, NEW = not in design
C = []
def s(key, screen, idt, en, note="", origin="DESIGN"):
    C.append((key, screen, origin, idt, en, note))

# 01 Pembuka
s("app_name", "Splash", "Izzah", "Izzah", "Brand name, same in both languages.")
s("splash_tagline", "Splash", "Menjaga diri, satu hari setiap kali", "Guarding yourself, one day at a time")
s("welcome_title", "Welcome", "Berdiri tegak lagi, satu hari setiap kali.", "Stand tall again, one day at a time.", "H1 42sp, line-height 1.04. EN is a bit longer, so check that it wraps to 3 lines max.")
s("welcome_quote", "Welcome", "“Barang siapa menghendaki izzah, maka milik Allah-lah izzah itu.” (QS 35:10)", "“Whoever desires honor, to Allah belongs all honor.” (QS 35:10)", "Qur'an translation. EN is based on Sahih International. Needs ustadz review.")
s("welcome_body", "Welcome", "Teman perjalanan untuk menjaga diri. Ada saat dorongan datang, ada saat kamu ingin mulai lagi.", "A companion for guarding yourself. Here when an urge comes, and here when you want to begin again.")
s("welcome_cta_start", "Welcome", "Mulai perjalanan", "Start the journey")
s("welcome_cta_pin", "Welcome", "Sudah pernah pakai? Masuk dengan PIN", "Used Izzah before? Sign in with PIN")
s("welcome_footnote", "Welcome", "Anonim. Tanpa email, tanpa nomor HP.", "Anonymous. No email, no phone number.")

# 02 Onboarding
s("common_continue", "Onboarding (all)", "Lanjut", "Continue")
s("common_back_cd", "All pushed screens", "Kembali", "Back", "Content description for the ← button (accessibility).", "NEW")
s("onboarding_progress", "Onboarding (all)", "{step}/6", "{step}/6", "Placeholder only.")
s("niat_title", "1 · Niat", "Apa yang membuatmu ingin berubah?", "What makes you want to change?")
s("niat_subtitle", "1 · Niat", "Pilih yang paling dekat di hati. Ini akan kami ingatkan saat kamu goyah.", "Pick what's closest to your heart. We'll remind you of it when you feel shaky.")
s("niat_opt_allah", "1 · Niat", "Dekat dengan Allah", "Closer to Allah")
s("niat_opt_family", "1 · Niat", "Untuk keluarga", "For my family")
s("niat_opt_spouse", "1 · Niat", "Pasangan / calon", "My spouse / future spouse")
s("niat_opt_focus", "1 · Niat", "Fokus & produktif", "Focus & productivity")
s("niat_opt_peace", "1 · Niat", "Ketenangan pikiran", "Peace of mind")
s("niat_opt_selfrespect", "1 · Niat", "Harga diri", "Self-respect")
s("niat_note", "1 · Niat", "Tersimpan hanya di HP-mu.", "Saved only on your phone.")
s("target_title", "2 · Target", "Pilih target pertamamu", "Choose your first goal")
s("target_subtitle", "2 · Target", "Kami menghitung hari yang kamu jaga. Terpeleset tidak menghapus hitungan.", "We count the days you keep. A slip doesn't erase the count.")
s("target_7_title", "2 · Target", "Satu minggu", "One week")
s("target_7_desc", "2 · Target", "Langkah pertama", "A first step")
s("target_30_title", "2 · Target", "Satu bulan", "One month")
s("target_30_desc", "2 · Target", "Membangun kebiasaan", "Building a habit")
s("target_40_title", "2 · Target", "Empat puluh hari", "Forty days")
s("target_40_desc", "2 · Target", "Waktu yang sering disebut untuk perubahan diri", "A span often linked to real self-change")
s("target_90_title", "2 · Target", "Tiga bulan", "Three months")
s("target_90_desc", "2 · Target", "Disarankan · otak mulai pulih", "Recommended · the brain begins to recover", "This is a health claim, so consider softening it, e.g. 'Recommended · time for real change'.")
s("rawan_title", "3 · Jam rawan", "Kapan biasanya paling berat?", "When is it usually hardest?")
s("rawan_subtitle", "3 · Jam rawan", "Kami akan lebih dekat di waktu-waktu itu: pengingat, dan pagar untuk aplikasi.", "We'll stay closer at those times, with reminders and a fence around certain apps.")
s("rawan_opt_late", "3 · Jam rawan", "Larut malam", "Late at night")
s("rawan_opt_dawn", "3 · Jam rawan", "Dini hari", "Before dawn")
s("rawan_opt_afternoon", "3 · Jam rawan", "Siang, sendirian", "Afternoons, alone")
s("rawan_opt_social", "3 · Jam rawan", "Setelah buka medsos", "After social media")
s("rawan_opt_stress", "3 · Jam rawan", "Saat stres", "When stressed")
s("rawan_opt_weekend", "3 · Jam rawan", "Akhir pekan", "Weekends")
s("izin_title", "4 · Izin", "Izinkan Izzah menjaga waktu rawanmu", "Let Izzah guard your hard hours")
s("izin_subtitle", "4 · Izin", "Dengan izin akses penggunaan, Izzah bisa memberi jeda pada aplikasi yang kamu pilih.", "With usage access, Izzah can put a pause in front of the apps you choose.", "Original: 'Dengan izin Screen Time…'. Screen Time is iOS-only.", "ANDROID")
s("izin_promise_1", "4 · Izin", "Riwayat browsing tidak pernah kami lihat atau simpan.", "We never see or store your browsing history.")
s("izin_promise_2", "4 · Izin", "Kamu sendiri yang memilih aplikasi dan jamnya.", "You choose the apps and the hours yourself.")
s("izin_promise_3", "4 · Izin", "Bisa dimatikan kapan saja dari Pengaturan.", "You can turn it off anytime in Settings.")
s("izin_cta", "4 · Izin", "Berikan izin", "Grant access", "Original: 'Izinkan Screen Time'. On Android it opens 2 system settings in turn.", "ANDROID")
s("izin_skip", "4 · Izin", "Nanti saja", "Maybe later")
s("izin_perm_usage_title", "4 · Izin", "Akses penggunaan", "Usage access", "Replaces the iOS system alert. Status row #1.", "NEW")
s("izin_perm_usage_desc", "4 · Izin", "Supaya Izzah tahu kapan aplikasi pilihanmu dibuka.", "So Izzah knows when an app you chose is opened.", "", "NEW")
s("izin_perm_overlay_title", "4 · Izin", "Tampil di atas aplikasi lain", "Display over other apps", "Status row #2.", "NEW")
s("izin_perm_overlay_desc", "4 · Izin", "Supaya layar jeda bisa muncul.", "So the pause screen can appear.", "", "NEW")
s("izin_perm_on", "4 · Izin", "Aktif", "On", "", "NEW")
s("izin_perm_off", "4 · Izin", "Belum", "Not yet", "", "NEW")
s("pin_create_title", "5 · PIN", "Buat PIN 4 angka", "Create a 4-digit PIN")
s("pin_create_subtitle", "5 · PIN", "Supaya hanya kamu yang bisa membuka Izzah.", "So only you can open Izzah.")
s("pin_confirm_title", "5 · PIN", "Ketik ulang PIN-mu", "Enter your PIN again", "Proposed confirmation step (not in design), so a typo doesn't lock the user out.", "NEW")
s("pin_mismatch", "5 · PIN", "PIN tidak sama. Coba lagi.", "PINs don't match. Try again.", "", "NEW")
s("pin_biometric_toggle", "5 · PIN", "Buka juga dengan sidik jari atau wajah", "Also unlock with fingerprint or face", "Original: 'Buka juga dengan Face ID'.", "ANDROID")
s("keypad_delete_cd", "5 · PIN / Unlock", "Hapus angka", "Delete digit", "Content description for ⌫.", "NEW")
s("notif_title", "6 · Notifikasi", "Pengingat yang tidak membuka rahasiamu", "Reminders that keep your secret")
s("notif_subtitle", "6 · Notifikasi", "Isi notifikasi selalu netral. Orang di sebelahmu tidak akan tahu.", "Notifications are always neutral. The person next to you won't know.")
s("notif_preview_time", "6 · Notifikasi", "sekarang", "now")
s("notif_preview_body", "6 · Notifikasi", "Waktunya Ashar. Sejenak untuk dirimu.", "Time for Asr. A moment for yourself.")
s("notif_opt_prayer", "6 · Notifikasi", "Pengingat sholat", "Prayer reminders")
s("notif_opt_prayer_desc", "6 · Notifikasi", "Lima waktu, isi netral", "All five prayers, neutral wording")
s("notif_opt_dzikir", "6 · Notifikasi", "Dzikir pagi & petang", "Morning & evening dhikr")
s("notif_opt_dzikir_desc", "6 · Notifikasi", "Setelah Subuh dan Ashar", "After Fajr and Asr")
s("notif_opt_checkin", "6 · Notifikasi", "Check-in malam", "Evening check-in")
s("notif_opt_checkin_desc", "6 · Notifikasi", "Pukul {time}, sebelum jam rawan", "At {time}, before the hard hours", "{time} = 21.30 (ID) / 9:30 PM (EN). Format with the locale, never hardcode it.")
s("notif_cta", "6 · Notifikasi", "Izinkan notifikasi", "Allow notifications")
s("done_day_label", "Selesai", "Hari pertama · {hijri}", "Day one · {hijri}")
s("done_journey", "Selesai", "Perjalanan {target} hari dimulai hari ini.", "Your {target}-day journey starts today.")
s("done_hint", "Selesai", "Kalau dorongan datang, tombol Jeda selalu ada di tengah bawah.", "If an urge comes, the Pause button is always at the bottom center.")
s("done_cta", "Selesai", "Masuk ke Beranda", "Go to Home")

# Tab bar
s("tab_home", "Tab bar", "Beranda", "Home")
s("tab_progress", "Tab bar", "Progres", "Progress")
s("tab_jeda", "Tab bar", "Jeda", "Pause", "Center button. Option: keep 'Jeda' in EN as a brand word.")
s("tab_protect", "Tab bar", "Lindungi", "Protect")
s("tab_settings", "Tab bar", "Pengaturan", "Settings")

# 03 Inti — Beranda
s("home_greeting", "Beranda", "Assalamu’alaikum", "Assalamu’alaikum", "Kept as-is in EN.")
s("home_date_line", "Beranda", "{weekday} · {hijri}", "{weekday} · {hijri}", "Example: 'Selasa · 13 Rabiul Akhir 1448 H' / 'Tuesday · 13 Rabi' al-Thani 1448 AH'.")
s("home_hero_label", "Beranda", "Kamu sudah terjaga", "You've kept", "Rendered UPPERCASE. Reads with the big number: 'YOU'VE KEPT 42 days'.")
s("home_hero_unit", "Beranda", "hari", "days", "PLURAL in EN (1 day, 2 days). ID has no plural.")
s("home_today_clock", "Beranda", "hari ini · {time}", "today · {time}")
s("home_shortcut_dzikir", "Beranda", "Dzikir", "Dhikr")
s("home_shortcut_sholat", "Beranda", "Sholat", "Prayer")
s("home_shortcut_jurnal", "Beranda", "Jurnal", "Journal")
s("home_shortcut_doa", "Beranda", "Doa", "Dua")
s("home_bars_caption", "Beranda", "{kept} dari {target} hari menuju target pertamamu", "{kept} of {target} days toward your first goal")
s("home_bars_hint", "Beranda", "Ketuk batang untuk detail · lihat semua →", "Tap a bar for details · see all →")
s("home_prayer_card", "Beranda", "Sholat hari ini · {done}/5", "Today's prayers · {done}/5", "Rendered UPPERCASE.")
s("home_next_prayer", "Beranda", "{prayer} · {time}", "{prayer} · {time}")
s("home_prayers_complete", "Beranda", "Alhamdulillah, lengkap", "Alhamdulillah, all done")
s("home_dzikir_tag", "Beranda", "Al-Ma’tsurat · 7 mnt", "Al-Ma’thurat · 7 min")
s("home_dzikir_card", "Beranda", "Dzikir petang hari ini", "Today's evening dhikr", "Becomes 'Dzikir pagi hari ini' / 'Today's morning dhikr' before Asr (proposed).")
s("home_dzikir_card_morning", "Beranda", "Dzikir pagi hari ini", "Today's morning dhikr", "Morning variant (proposed).", "NEW")
s("home_slip_prompt", "Beranda", "Tadi sempat terpeleset? Pintu taubat selalu terbuka.", "Slipped earlier? The door of repentance is always open.")
s("home_slip_cta", "Beranda", "Mulai lagi →", "Begin again →")
s("avatar_cd", "Beranda", "Buka pengaturan", "Open settings", "Content description for the avatar.", "NEW")

# Day sheet
s("day_title", "Detail hari", "Hari ke-{n}", "Day {n}", "Rendered UPPERCASE in the sheet.")
s("day_status_kept", "Detail hari / Progres", "Terjaga", "Kept")
s("day_status_slip", "Detail hari / Progres", "Terpeleset", "Slipped", "Never use red or 'reset' language for this state.")
s("day_status_today", "Detail hari / Progres", "Hari ini", "Today")
s("day_checkin_yes", "Detail hari", "Check-in · tercatat", "Check-in · logged")
s("day_checkin_no", "Detail hari", "Tidak ada check-in", "No check-in")
s("day_note_slip", "Detail hari", "Terpeleset malam itu, lalu kamu menekan “Mulai lagi” keesokan paginya.", "You slipped that night, then tapped “Begin again” the next morning.", "Auto note for a slip day with no journal entry.")
s("day_note_today", "Detail hari", "Hari ini masih berjalan.", "Today is still in progress.")
s("day_note_empty", "Detail hari", "Tidak ada catatan.", "No note.", "Kept day with no journal entry.", "NEW")
s("common_close", "Detail hari", "Tutup", "Close")

# Jeda
s("jeda_header", "Jeda", "Jeda · {seconds} detik", "Pause · {seconds} seconds", "Rendered UPPERCASE, countdown 60→0. PLURAL in EN.")
s("jeda_reminder", "Jeda", "Ingat {kept} hari yang sudah kamu jaga. Ini juga akan lewat.", "Remember the {kept} days you've kept. This too will pass.")
s("jeda_inhale", "Jeda", "Tarik napas…", "Breathe in…")
s("jeda_exhale", "Jeda", "Hembuskan…", "Breathe out…")
s("jeda_translit", "Jeda", "Astaghfirullāh al-‘aẓīm · napas ke-{n}", "Astaghfirullāh al-‘aẓīm · breath {n}")
s("jeda_calmer", "Jeda", "Aku sudah lebih tenang", "I feel calmer now")
s("jeda_read_doa", "Jeda", "Baca doa perlindungan", "Read a dua for protection")

# Progres
s("progress_title", "Progres", "Progres", "Progress")
s("progress_of_target", "Progres", "/ {target} hari", "/ {target} days")
s("progress_summary", "Progres", "Hari ke-{day} sejak mulai · {slips} hari terpeleset, semuanya kamu lanjutkan.", "Day {day} since you started · {slips} slips, and you kept going after every one.", "PLURAL in EN for {slips}. When slips = 0, use the next row.")
s("progress_summary_noslip", "Progres", "Hari ke-{day} sejak mulai.", "Day {day} since you started.", "", "NEW")
s("progress_legend_future", "Progres", "Belum dijalani", "Not yet")
s("progress_bar_caption", "Progres", "Makin tinggi batang, makin banyak sholat dan check-in di hari itu.", "The taller the bar, the more prayers and check-ins that day.")
s("progress_sel_title", "Progres", "Hari ke-{n} · {date}", "Day {n} · {date}")
s("progress_sel_sub", "Progres", "{hijri} · sholat {done}/5", "{hijri} · prayers {done}/5")
s("progress_stat_prayer", "Progres", "sholat terjaga", "prayers kept")
s("progress_stat_checkin", "Progres", "hari check-in", "check-in days")
s("progress_stat_jeda", "Progres", "jeda berhasil", "pauses completed")
s("progress_insight_title", "Progres", "Pola yang terlihat", "A pattern we noticed")
s("progress_insight_body", "Progres", "Kamu paling sering butuh jeda antara pukul {start} dan {end}. Mau diingatkan dzikir sebelum tidur?", "You most often need a pause between {start} and {end}. Want a dhikr reminder before bed?", "Show only when there is enough jeda data (proposed: ≥5 pauses).")
s("progress_remind_cta", "Progres", "Ingatkan jam {time}", "Remind me at {time}")
s("progress_remind_on", "Progres", "✓ Diingatkan jam {time}", "✓ Reminder set for {time}")

# Sholat
s("sholat_title", "Sholat", "Sholat hari ini", "Today's prayers")
s("sholat_of_five", "Sholat", "/ 5", "/ 5")
s("sholat_location", "Sholat", "{city} · {hijri}", "{city} · {hijri}")
s("sholat_footer", "Sholat", "Check-in setelah sholat ikut membuat batang hari ini lebih tinggi.", "Checking in after each prayer makes today's bar taller.")
s("sholat_location_pick", "Sholat", "Pilih kota", "Choose city", "Fallback when location is off or denied.", "NEW")
s("sholat_use_location", "Sholat", "Pakai lokasi sekarang", "Use current location", "", "NEW")
s("prayer_fajr", "Prayer names", "Subuh", "Fajr")
s("prayer_dhuhr", "Prayer names", "Dzuhur", "Dhuhr")
s("prayer_asr", "Prayer names", "Ashar", "Asr")
s("prayer_maghrib", "Prayer names", "Maghrib", "Maghrib")
s("prayer_isha", "Prayer names", "Isya", "Isha")

# Dzikir
s("dzikir_title", "Dzikir", "Dzikir setelah sholat", "Dhikr after prayer")
s("dzikir_hint", "Dzikir", "dari 33 · ketuk", "of 33 · tap")
s("dzikir_done", "Dzikir", "selesai · 99", "done · 99")
s("dzikir_footer", "Dzikir", "Getar halus setiap 33. Bisa juga pakai tombol volume.", "A gentle vibration every 33. You can also use the volume buttons.")
s("dzikir_reset_cd", "Dzikir", "Ulangi dari awal", "Start over", "Content description for ↺.", "NEW")

# Doa
s("doa_title", "Doa", "Doa saat goyah", "Duas for shaky moments")
s("doa_subtitle", "Doa", "Baca pelan-pelan. Tidak perlu hafal, cukup hadirkan hati.", "Read slowly. No need to memorize. Just bring your heart.")

# Jurnal
s("jurnal_title", "Jurnal", "Jurnal", "Journal")
s("jurnal_question", "Jurnal", "Bagaimana hatimu hari ini?", "How is your heart today?")
s("mood_strong", "Jurnal", "Kuat", "Strong")
s("mood_okay", "Jurnal", "Biasa", "Okay")
s("mood_shaky", "Jurnal", "Goyah", "Shaky")
s("mood_heavy", "Jurnal", "Berat", "Heavy")
s("jurnal_placeholder", "Jurnal", "Tulis sedikit. Apa yang membantu, apa yang berat…", "Write a little. What helped, what felt heavy…")
s("jurnal_locked_note", "Jurnal", "Terkunci dengan PIN. Hanya kamu yang bisa membaca.", "Locked with your PIN. Only you can read it.")
s("common_save", "Jurnal", "Simpan", "Save")
s("jurnal_saved_toast", "Jurnal", "Tersimpan. Terima kasih sudah jujur pada dirimu.", "Saved. Thank you for being honest with yourself.")
s("jurnal_earlier", "Jurnal", "Sebelumnya", "Earlier")
s("jurnal_entry_meta", "Jurnal", "{weekday} · {mood}", "{weekday} · {mood}")
s("jurnal_empty", "Jurnal", "Belum ada catatan. Tulisan pertamamu akan muncul di sini.", "No entries yet. Your first one will show up here.", "Empty state (design shows sample entries only).", "NEW")

# Mulai lagi
s("mulai_title", "Mulai lagi", "Terpeleset bukan akhir perjalanan.", "A slip is not the end of the journey.")
s("mulai_body", "Mulai lagi", "{kept} hari yang kamu jaga tetap tercatat. Hari ini cukup jadi satu titik redup di antaranya.", "The {kept} days you've kept still count. Today is just one dim dot among them.")
s("mulai_step_1", "Mulai lagi", "Istighfar", "Istighfar")
s("mulai_step_1_desc", "Mulai lagi", "Astaghfirullāh, 3×", "Astaghfirullāh, 3×")
s("mulai_step_2", "Mulai lagi", "Wudhu & sholat taubat", "Wudu & the prayer of repentance")
s("mulai_step_2_desc", "Mulai lagi", "2 rakaat, kapan pun kamu siap", "2 rak‘ahs, whenever you're ready")
s("mulai_step_3", "Mulai lagi", "Satu langkah kecil", "One small step")
s("mulai_step_3_desc", "Mulai lagi", "Hapus satu pemicu, misalnya satu aplikasi", "Remove one trigger, like a single app")
s("mulai_triggers_title", "Mulai lagi", "Apa pemicunya? (hanya untukmu)", "What triggered it? (just for you)")
s("trigger_alone_night", "Mulai lagi", "Sendirian malam", "Alone at night")
s("trigger_bored", "Mulai lagi", "Bosan", "Bored")
s("trigger_work_stress", "Mulai lagi", "Stres kerja", "Work stress")
s("trigger_social", "Mulai lagi", "Media sosial", "Social media")
s("trigger_tired", "Mulai lagi", "Lelah", "Tired")
s("trigger_no_sleep", "Mulai lagi", "Sulit tidur", "Couldn't sleep")
s("mulai_cta", "Mulai lagi", "Mulai lagi hari ini", "Begin again today")
s("mulai_close_cd", "Mulai lagi", "Tutup", "Close", "Content description for ✕.", "NEW")

# 04 Lindungi
s("protect_title", "Lindungi", "Lindungi", "Protect")
s("protect_desc", "Lindungi", "Aplikasi tertentu dikunci di jam-jam rawan. Bukan hukuman, hanya pagar kecil supaya kamu tidak berjuang sendirian.", "Chosen apps are locked during your hard hours. Not a punishment, just a small fence so you don't fight alone.")
s("protect_schedule", "Lindungi", "Jadwal", "Schedule")
s("slot_tag_risky", "Lindungi", "Rawan", "High-risk")
s("slot_tag_optional", "Lindungi", "Opsional", "Optional")
s("slot_tag_auto", "Lindungi", "Otomatis", "Automatic")
s("slot_1_desc", "Lindungi", "Setiap hari · dari pola jedamu", "Every day · based on your pauses")
s("slot_2_desc", "Lindungi", "Senin–Jumat · jam sepi di kantor", "Mon–Fri · quiet hours at work")
s("slot_3_title", "Lindungi", "Saat sendirian di rumah", "When alone at home", "Needs location permission + geofencing. Proposal: leave out of v1.")
s("slot_3_desc", "Lindungi", "Saat lokasi Rumah & tidak ada jadwal", "When you're at Home & nothing is scheduled")
s("protect_locked_apps", "Lindungi", "Yang dikunci · {n} aplikasi", "Locked · {n} apps", "PLURAL in EN (1 app).")
s("protect_add_app", "Lindungi", "+ Tambah aplikasi", "+ Add app", "Opens the installed-app picker (replaces iOS FamilyActivityPicker).", "NEW")
s("protect_when_opened", "Lindungi", "Saat dibuka", "When opened")
s("mode_gentle", "Lindungi", "Pengingat lembut", "Gentle reminder")
s("mode_gentle_desc", "Lindungi", "Muncul pesan, tetap bisa lanjut", "A message appears, you can still continue")
s("mode_pause", "Lindungi", "Jeda 60 detik dulu", "60-second pause first")
s("mode_pause_desc", "Lindungi", "Napas dulu, lalu boleh buka 5 menit", "Breathe first, then it opens for 5 minutes")
s("mode_full", "Lindungi", "Kunci penuh", "Full lock")
s("mode_full_desc", "Lindungi", "Tidak bisa dibuka sampai jadwal selesai", "Can't be opened until the schedule ends")
s("protect_preview_shield", "Lindungi", "Lihat tampilan saat terkunci", "Preview the locked screen")

# Terkunci
s("shield_label", "Layar terkunci", "Waktu terjaga · sampai {time}", "Guarded hours · until {time}", "Rendered UPPERCASE.")
s("shield_title", "Layar terkunci", "Sudah larut. Istirahat dulu, yuk.", "It's late. Time to rest, okay?", "Fits a late-night slot only. A daytime variant is proposed below.")
s("shield_title_day", "Layar terkunci", "Sebentar dulu. Tarik napas, yuk.", "Hold on a moment. Let's take a breath.", "For daytime slots (e.g. 13.00–15.00).", "NEW")
s("shield_body", "Layar terkunci", "Kamu sendiri yang memasang pagar ini, di saat kamu sedang kuat.", "You put up this fence yourself, at a moment you were strong.")
s("shield_take_jeda", "Layar terkunci", "Ambil jeda 60 detik", "Take a 60-second pause")
s("shield_close_app", "Layar terkunci", "Tutup aplikasi", "Close the app")
s("shield_continue", "Layar terkunci", "Tetap lanjutkan", "Continue anyway", "Needed for 'Pengingat lembut' mode, where the design has no continue button.", "NEW")
s("shield_note_gentle", "Layar terkunci", "Kamu tetap bisa melanjutkan jika memang perlu", "You can still continue if you really need to")
s("shield_note_pause", "Layar terkunci", "Setelah jeda, aplikasi terbuka 5 menit", "After the pause, the app opens for 5 minutes")
s("shield_note_full", "Layar terkunci", "Terbuka lagi pukul {time}", "Opens again at {time}")

# Pengaturan
s("settings_title", "Pengaturan", "Pengaturan", "Settings")
s("settings_profile_title", "Pengaturan", "Akun anonim", "Anonymous account")
s("settings_profile_sub", "Pengaturan", "Tanpa email · data hanya di HP ini", "No email · data stays on this phone")
s("settings_group_privacy", "Pengaturan", "Privasi", "Privacy")
s("settings_pin_lock", "Pengaturan", "Kunci dengan PIN & sidik jari", "Lock with PIN & fingerprint", "Original: '…& Face ID'.", "ANDROID")
s("settings_disguise", "Pengaturan", "Ikon & nama samaran", "Disguised icon & name")
s("settings_hide_search", "Pengaturan", "Sembunyikan dari pencarian HP", "Hide from phone search", "Android has no reliable API for this. Proposal: remove it, or replace it with 'Sembunyikan isi di layar Recent' / 'Hide content in Recents' (FLAG_SECURE).")
s("settings_hide_recents", "Pengaturan", "Sembunyikan isi di layar Recent", "Hide content in Recents", "Android replacement for the row above (FLAG_SECURE also blocks screenshots).", "NEW")
s("settings_group_journey", "Pengaturan", "Perjalanan", "Journey")
s("settings_target", "Pengaturan", "Target", "Goal")
s("settings_target_value", "Pengaturan", "{n} hari", "{n} days", "PLURAL in EN.")
s("settings_reminders", "Pengaturan", "Pengingat sholat & dzikir", "Prayer & dhikr reminders")
s("settings_schedule", "Pengaturan", "Jadwal waktu terjaga", "Guarded-hours schedule")
s("settings_group_data", "Pengaturan", "Data & bantuan", "Data & help")
s("settings_export_delete", "Pengaturan", "Ekspor & hapus data", "Export & delete data")
s("settings_counselor", "Pengaturan", "Bicara dengan konselor", "Talk to a counselor", "Where does it go? Needs a real contact or link before release.")
s("settings_about", "Pengaturan", "Tentang Izzah", "About Izzah")
s("settings_language", "Pengaturan", "Bahasa", "Language", "Language picker row, under Privasi or a new 'Umum' group.", "NEW")
s("language_system", "Pengaturan", "Ikuti sistem", "System default", "", "NEW")
s("language_id", "Pengaturan", "Bahasa Indonesia", "Bahasa Indonesia", "Always shown in its own language.", "NEW")
s("language_en", "Pengaturan", "English", "English", "Always shown in its own language.", "NEW")

# Ikon samaran
s("disguise_title", "Ikon samaran", "Ikon samaran", "Disguised icon")
s("disguise_desc", "Ikon samaran", "Ganti ikon dan nama di layar HP supaya tidak menarik perhatian.", "Change the icon and name on your home screen so it doesn't draw attention.")
s("disguise_izzah", "Ikon samaran", "Izzah", "Izzah")
s("disguise_steps", "Ikon samaran", "Langkah", "Steps", "This becomes the launcher label, so it's localized too.")
s("disguise_focus", "Ikon samaran", "Fokus", "Focus")
s("disguise_notes", "Ikon samaran", "Catatan", "Notes")
s("disguise_footer", "Ikon samaran", "Nama di notifikasi tetap “Izzah”, tapi isinya selalu netral.", "Notifications still show “Izzah”, but their content is always neutral.", "Original: 'Notifikasi akan memakai nama yang sama.' Android always shows the application label in notifications; an activity-alias can't change it.", "ANDROID")

# Data
s("data_title", "Datamu", "Datamu", "Your data")
s("data_desc", "Datamu", "Semua tersimpan di HP ini. Tidak ada di server kami.", "Everything is stored on this phone. Nothing is on our servers.")
s("data_export_title", "Datamu", "Ekspor cadangan", "Export a backup")
s("data_export_desc", "Datamu", "File terenkripsi dengan PIN-mu. Pakai saat pindah HP.", "A file encrypted with your PIN. Use it when you switch phones.")
s("data_export_cta", "Datamu", "Buat file cadangan", "Create backup file")
s("data_import_wrong_pin", "Datamu", "PIN cadangan tidak cocok.", "The backup PIN doesn't match.", "", "NEW")
s("data_export_done", "Datamu", "Cadangan tersimpan.", "Backup saved.", "Toast after the system save dialog.", "NEW")
s("data_import_cta", "Datamu", "Pulihkan dari cadangan", "Restore from backup", "Without this, the export is useless on a new phone. Proposed for v1.", "NEW")
s("data_delete_title", "Datamu", "Hapus semua data", "Delete all data")
s("data_delete_desc", "Datamu", "Progres, jurnal, dan pengaturan akan hilang permanen.", "Your progress, journal and settings will be permanently erased.")
s("data_delete_cta", "Datamu", "Hapus data", "Delete data")
s("delete_sheet_title", "Hapus (sheet)", "Yakin hapus semua?", "Delete everything?")
s("delete_sheet_body", "Hapus (sheet)", "Ini tidak bisa dibatalkan. Kalau hanya ingin istirahat, kamu bisa menjeda pengingat saja.", "This can't be undone. If you just need a break, you can pause the reminders instead.")
s("delete_sheet_confirm", "Hapus (sheet)", "Ya, hapus semuanya", "Yes, delete everything")
s("common_cancel", "Hapus (sheet)", "Batal", "Cancel")

# Unlock
s("unlock_title", "Buka dengan PIN", "Masukkan PIN", "Enter your PIN")
s("unlock_subtitle", "Buka dengan PIN", "Assalamu’alaikum. Selamat datang kembali.", "Assalamu’alaikum. Welcome back.")
s("unlock_wrong", "Buka dengan PIN", "PIN salah. Coba lagi.", "Wrong PIN. Try again.", "", "NEW")
s("unlock_cooldown", "Buka dengan PIN", "Terlalu banyak percobaan. Coba lagi dalam {seconds} detik.", "Too many tries. Try again in {seconds} seconds.", "After 5 wrong PINs: 30 s cooldown. PLURAL in EN.", "NEW")
s("unlock_forgot", "Buka dengan PIN", "Lupa PIN?", "Forgot your PIN?", "", "NEW")
s("unlock_forgot_body", "Buka dengan PIN", "Karena akunmu anonim, PIN tidak bisa dipulihkan. Kamu bisa menghapus data dan mulai dari awal, atau memulihkan dari file cadangan.", "Your account is anonymous, so the PIN can't be recovered. You can delete your data and start fresh, or restore from a backup file.", "Bottom sheet. Buttons reuse delete_sheet_confirm + data_import_cta + common_cancel.", "NEW")
s("biometric_title", "Buka dengan PIN", "Buka Izzah", "Unlock Izzah", "Title of the system BiometricPrompt. Shows the disguise name if a disguise is active.", "NEW")
s("biometric_use_pin", "Buka dengan PIN", "Pakai PIN", "Use PIN", "", "NEW")

# Notifications (neutral)
s("notif_prayer_body", "Notifikasi (sistem)", "Waktunya {prayer}. Sejenak untuk dirimu.", "Time for {prayer}. A moment for yourself.", "RULE: no notification may mention porn, relapse or addiction.")
s("notif_dzikir_morning", "Notifikasi (sistem)", "Pagi yang baru. Ada waktu sebentar untuk dirimu?", "A fresh morning. Got a minute for yourself?", "", "NEW")
s("notif_dzikir_evening", "Notifikasi (sistem)", "Sore menjelang. Sejenak untuk menenangkan hati.", "Evening is here. A moment to settle your heart.", "", "NEW")
s("notif_checkin", "Notifikasi (sistem)", "Bagaimana harimu? Catat sebentar sebelum istirahat.", "How was your day? A quick note before you rest.", "", "NEW")
s("notif_channel_reminders", "Notifikasi (sistem)", "Pengingat", "Reminders", "Android notification channel name (visible in system settings).", "NEW")
s("notif_channel_protect", "Notifikasi (sistem)", "Penjagaan aktif", "Protection running", "Foreground service channel for the blocker.", "NEW")
s("notif_protect_ongoing", "Notifikasi (sistem)", "Berjalan di latar belakang.", "Running in the background.", "Persistent foreground-service notification. Must stay neutral.", "NEW")

# Hijri
hijri = [("Muharram","Muharram"),("Safar","Safar"),("Rabiul Awal","Rabi‘ al-Awwal"),("Rabiul Akhir","Rabi‘ al-Thani"),
         ("Jumadil Awal","Jumada al-Ula"),("Jumadil Akhir","Jumada al-Akhirah"),("Rajab","Rajab"),("Sya’ban","Sha‘ban"),
         ("Ramadhan","Ramadan"),("Syawal","Shawwal"),("Dzulqa’dah","Dhu al-Qa‘dah"),("Dzulhijjah","Dhu al-Hijjah")]
for i,(a,b) in enumerate(hijri,1):
    s(f"hijri_month_{i}", "Hijri calendar", a, b, "Used by {hijri}." if i==1 else "")
s("hijri_suffix", "Hijri calendar", "{day} {month} {year} H", "{day} {month} {year} AH")

RELIGIOUS = [
    ("app_name_arabic", "Splash", "عِزَّة", "—", "—", "—", "Logo subtitle"),
    ("done_bismillah", "Selesai", "بِسْمِ ٱللَّٰهِ", "Bismillāh", "Dengan nama Allah", "In the name of Allah", ""),
    ("home_dzikir_arabic", "Beranda", "حَسْبُنَا ٱللَّٰهُ", "Ḥasbunallāh", "Cukuplah Allah bagi kami", "Allah is sufficient for us", "Translation not shown in UI (Arabic only)."),
    ("jeda_istighfar", "Jeda", "أَسْتَغْفِرُ ٱللَّٰهَ ٱلْعَظِيمَ", "Astaghfirullāh al-‘aẓīm", "Aku memohon ampun kepada Allah Yang Maha Agung", "I seek forgiveness from Allah, the Magnificent", "Translation not shown in UI."),
    ("dzikir_1", "Dzikir", "سُبْحَانَ ٱللَّٰهِ", "Subḥānallāh", "Maha Suci Allah", "Glory be to Allah", ""),
    ("dzikir_2", "Dzikir", "ٱلْحَمْدُ لِلَّٰهِ", "Alḥamdulillāh", "Segala puji bagi Allah", "All praise is for Allah", ""),
    ("dzikir_3", "Dzikir", "ٱللَّٰهُ أَكْبَرُ", "Allāhu akbar", "Allah Maha Besar", "Allah is the Greatest", ""),
    ("doa_1", "Doa", "أَعُوذُ بِٱللَّٰهِ مِنَ ٱلشَّيْطَانِ ٱلرَّجِيمِ", "A‘ūdzu billāhi minasy-syaiṭānir-rajīm",
     "“Aku berlindung kepada Allah dari setan yang terkutuk.”", "“I seek refuge in Allah from Satan, the accursed.”", "EN transliteration could use 'A‘ūdhu billāhi mina-sh-shayṭāni-r-rajīm'."),
    ("doa_2", "Doa", "ٱللَّٰهُمَّ إِنِّي أَعُوذُ بِكَ مِنْ شَرِّ سَمْعِي، وَمِنْ شَرِّ بَصَرِي، وَمِنْ شَرِّ لِسَانِي، وَمِنْ شَرِّ قَلْبِي، وَمِنْ شَرِّ مَنِيِّي",
     "Allāhumma innī a‘ūdzu bika min syarri sam‘ī, wa min syarri baṣarī, wa min syarri lisānī, wa min syarri qalbī, wa min syarri maniyyī",
     "“Ya Allah, aku berlindung kepada-Mu dari keburukan pendengaranku, penglihatanku, lisanku, hatiku, dan syahwatku.” (HR Abu Dawud, Tirmidzi)",
     "“O Allah, I seek refuge in You from the evil of my hearing, my sight, my tongue, my heart, and my desire.” (Abu Dawud, Tirmidhi)",
     "Transliteration uses the Indonesian convention (sy, dz). For EN, consider 'sh', 'dh'."),
    ("doa_3", "Doa", "يَا مُقَلِّبَ ٱلْقُلُوبِ ثَبِّتْ قَلْبِي عَلَىٰ دِينِكَ", "Yā muqallibal-qulūb, tsabbit qalbī ‘alā dīnik",
     "“Wahai Yang membolak-balikkan hati, teguhkan hatiku di atas agama-Mu.” (HR Tirmidzi)",
     "“O Turner of hearts, keep my heart firm upon Your religion.” (Tirmidhi)", "For EN, 'thabbit' is the more common transliteration."),
    ("welcome_quote", "Welcome", "—", "—", "“Barang siapa menghendaki izzah, maka milik Allah-lah izzah itu.” (QS 35:10)",
     "“Whoever desires honor, to Allah belongs all honor.” (QS 35:10)", "Also in the Copy sheet. Qur'an translation."),
]

GLOSSARY = [
    ("Terjaga (hari terjaga)", "Kept (days kept)", "The core idea: we count days KEPT, never a 'streak'. Avoid 'clean', 'sober' and 'streak' in EN."),
    ("Terpeleset", "Slip / Slipped", "Never 'relapse', 'fail' or 'reset'. No red color."),
    ("Jeda", "Pause", "Alternative: keep 'Jeda' as a brand word in EN too."),
    ("Mulai lagi", "Begin again", "Not 'restart' or 'reset', which imply losing progress."),
    ("Lindungi", "Protect", ""),
    ("Waktu terjaga / jam rawan", "Guarded hours / hard hours", "'Rawan' as a tag = 'High-risk'."),
    ("Pagar", "Fence", "Metaphor for app locking. Kept literal on purpose."),
    ("Goyah", "Shaky", ""),
    ("Sholat / Dzikir / Doa", "Prayer / Dhikr / Dua", "Prayer names in EN: Fajr, Dhuhr, Asr, Maghrib, Isha."),
    ("Taubat", "Repentance", ""),
    ("Kamu (informal)", "You (warm, plain)", "Tone: warm, short, a friend who doesn't judge. Use contractions in EN."),
    ("HP", "Phone", ""),
]

wb = Workbook()
F = "Arial"
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
hdr_fill = PatternFill("solid", fgColor="2A2A30")
review_fill = PatternFill("solid", fgColor="FFF2CC")
origin_fill = {"DESIGN": None, "ANDROID": PatternFill("solid", fgColor="DDEBF7"), "NEW": PatternFill("solid", fgColor="E2EFDA")}
body = Font(name=F, size=10)
wrap = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="D9D9D9")
border = Border(bottom=thin)

def table(ws, headers, rows, widths, review_cols=(), status_col=None):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.alignment = hdr_font, hdr_fill, Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment, c.border = body, wrap, border
            if c.column in review_cols:
                c.fill = review_fill
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 22
    if status_col:
        dv = DataValidation(type="list", formula1='"OK,Revisi,Hapus"', allow_blank=True)
        ws.add_data_validation(dv)
        L = ws.cell(1, status_col).column_letter
        dv.add(f"{L}2:{L}{ws.max_row}")

ws = wb.active
ws.title = "README"
readme = [
    ("Izzah · Copy deck (ID + EN)", True),
    ("Source: design_handoff_izzah_app/Izzah App.dc.html (prototype). Target: Android Jetpack Compose.", False),
    ("", False),
    ("How to review", True),
    ("1. Open the 'Copy' sheet. Read the Bahasa Indonesia and English columns.", False),
    ("2. Set the Status column (yellow): OK / Revisi / Hapus. Write your fix in 'Komentar / revisi'.", False),
    ("3. The 'Religious' sheet holds Arabic, transliterations and Qur'an/hadith meanings. It must be checked by an ustadz before release.", False),
    ("4. The 'Glossary' sheet sets the key terms. Settle it first so the whole deck stays consistent.", False),
    ("", False),
    ("Origin column", True),
    ("DESIGN = verbatim from the prototype (only translated to EN).", False),
    ("ANDROID (blue) = changed from an iOS-only string (Screen Time, Face ID). Check the wording.", False),
    ("NEW (green) = not in the design. Needed for Android, accessibility, or a missing state. Approve or delete.", False),
    ("", False),
    ("Placeholders", True),
    ("{kept}, {target}, {time}, {hijri} and similar are filled by the app. Don't rename them. Times and dates are formatted per locale (21.30 vs 9:30 PM).", False),
    ("Rows marked PLURAL need singular/plural forms in EN (1 day / 2 days).", False),
    ("", False),
    ("Example of a filled review row", True),
    ("Key: tab_jeda · EN: Pause · Status: Revisi · Komentar: keep 'Jeda' as a brand word", False),
]
for text, bold in readme:
    ws.append([text])
    ws.cell(ws.max_row, 1).font = Font(name=F, size=13 if bold and ws.max_row == 1 else 10, bold=bold)
ws.column_dimensions["A"].width = 120

ws = wb.create_sheet("Copy")
rows = [(i + 1, k, sc, o, idt, en, n, "", "") for i, (k, sc, o, idt, en, n) in enumerate(C)]
table(ws, ["#", "Android key", "Screen", "Origin", "Bahasa Indonesia", "English", "Catatan", "Status", "Komentar / revisi"],
      rows, [5, 26, 18, 10, 48, 48, 44, 10, 36], review_cols=(8, 9), status_col=8)
for row in ws.iter_rows(min_row=2):
    f = origin_fill[row[3].value]
    if f:
        for c in row[:7]:
            c.fill = f

ws = wb.create_sheet("Religious")
table(ws, ["Key", "Screen", "Arabic", "Transliteration", "Arti (ID)", "Meaning (EN)", "Catatan", "Status", "Komentar ustadz"],
      RELIGIOUS, [18, 12, 40, 38, 44, 44, 34, 10, 34], review_cols=(8, 9), status_col=8)
for row in ws.iter_rows(min_row=2):
    row[2].font = Font(name="Arial", size=14)
    row[2].alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")

ws = wb.create_sheet("Glossary")
table(ws, ["Bahasa Indonesia", "English", "Aturan / alasan", "Status", "Komentar"],
      GLOSSARY, [30, 30, 70, 10, 36], review_cols=(4, 5), status_col=4)

wb.save(OUT)
from collections import Counter
print(len(C), "copy rows", Counter(o for _, _, o, *_ in C), len(RELIGIOUS), "religious", len(GLOSSARY), "glossary")
