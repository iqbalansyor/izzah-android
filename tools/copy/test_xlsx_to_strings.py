import importlib.util, pathlib
spec = importlib.util.spec_from_file_location("x", pathlib.Path(__file__).with_name("xlsx_to_strings.py"))
x = importlib.util.module_from_spec(spec); spec.loader.exec_module(x)

def test_placeholders_positional_and_typed():
    assert x.convert("{kept} dari {target} hari") == "%1$d dari %2$d hari"
    assert x.convert("{prayer} · {time}") == "%1$s · %2$s"
    assert x.convert("{n} of {n}") == "%1$d of %1$d"

def test_escaping():
    assert x.convert("You've kept") == "You\\'ve kept"
    assert x.convert('He said "hi"') == 'He said \\"hi\\"'
    assert x.convert("@home") == "\\@home"
    assert x.convert("?x") == "\\?x"
    assert x.convert("100%") == "100%"          # plain string: % stays literal
    assert x.convert("{n}%") == "%1$d%%"        # formatted string: % must be doubled

def test_english_singular():
    assert x.singular("{n} days") == "{n} day"
    assert x.singular("Locked · {n} apps") == "Locked · {n} app"
    assert x.singular("Pause · {seconds} seconds") == "Pause · {seconds} second"
    assert x.singular("days") == "day"
    assert x.singular("{slips} slips") == "{slips} slip"

def test_hapus_rows_skipped_and_plurals(tmp_path):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Copy"
    ws.append(["#","Android key","Screen","Origin","Bahasa Indonesia","English","Catatan","Status","Komentar / revisi"])
    ws.append([1,"a","s","DESIGN","Halo","Hello","","",""])
    ws.append([2,"b","s","NEW","Hapus aku","Drop me","","Hapus",""])
    ws.append([3,"c","s","DESIGN","{n} hari","{n} days","PLURAL in EN.","OK",""])
    r = wb.create_sheet("Religious")
    r.append(["Key","Screen","Arabic","Transliteration","Arti (ID)","Meaning (EN)","Catatan","Status","Komentar ustadz"])
    r.append(["dzikir_1","D","سُبْحَانَ","Subḥānallāh","Maha Suci Allah","Glory be to Allah","","",""])
    p = tmp_path / "c.xlsx"; wb.save(p)
    en, idn = x.build(p)
    assert '<string name="a">Hello</string>' in en and '<string name="a">Halo</string>' in idn
    assert 'name="b"' not in en
    assert '<item quantity="one">%1$d day</item>' in en and '<item quantity="other">%1$d days</item>' in en
    assert '<item quantity="one">%1$d hari</item>' in idn
    assert 'name="dzikir_1" translatable="false"' in en and 'name="dzikir_1"' not in idn
    assert '<string name="dzikir_1_meaning">Maha Suci Allah</string>' in idn
